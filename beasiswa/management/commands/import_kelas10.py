import os
import openpyxl
from django.core.management.base import BaseCommand
from django.db import transaction
from beasiswa.models import Siswa, PenilaianBeasiswa, DataTraining, HasilSeleksi
from beasiswa.utils.c45 import run_seleksi


class Command(BaseCommand):
    help = "Import data siswa dan penilaian kelas X dari file Excel (f_leger_X B.xlsx)"

    def add_arguments(self, parser):
        parser.add_argument(
            '--file',
            type=str,
            default='f_leger_X B.xlsx',
            help='Path ke file Excel leger nilai kelas X'
        )
        parser.add_argument(
            '--run-seleksi',
            action='store_true',
            help='Otomatis jalankan klasifikasi C4.5 setelah import'
        )

    def handle(self, *args, **options):
        file_path = options['file']
        if not os.path.exists(file_path):
            self.stderr.write(self.style.ERROR(f"File tidak ditemukan: {file_path}"))
            return

        self.stdout.write(self.style.NOTICE(f"Membaca file: {file_path}..."))

        wb = openpyxl.load_workbook(file_path, data_only=True)
        if 'Worksheet' not in wb.sheetnames or 'Sheet2' not in wb.sheetnames:
            self.stderr.write(self.style.ERROR("File Excel harus memiliki sheet 'Worksheet' dan 'Sheet2'"))
            return

        ws = wb['Worksheet']
        s2 = wb['Sheet2']

        # 1. Parse Sheet2 (Data Sosial & Ekonomi)
        sheet2_data = {}
        for r in range(2, s2.max_row + 1):
            nama_cell = s2.cell(row=r, column=2).value
            if not nama_cell:
                continue
            nama = str(nama_cell).strip().upper()
            sktm_val = str(s2.cell(row=r, column=5).value or '').strip().upper()
            penghasilan_val = s2.cell(row=r, column=6).value or 0
            prestasi_val = str(s2.cell(row=r, column=7).value or 'RENDAH').strip().title()
            jto_val = s2.cell(row=r, column=8).value or 1

            sheet2_data[nama] = {
                'sktm': sktm_val in ['IYA', 'YA', 'TRUE', '1'],
                'penghasilan': int(penghasilan_val),
                'prestasi': prestasi_val if prestasi_val in ['Tinggi', 'Sedang', 'Rendah'] else 'Rendah',
                'jto': int(jto_val),
            }

        # Mapping gender perempuan untuk kelas X B
        FEMALE_NAMES = {
            'ALIS MAULIYA', 'ALIS MAULIA', 'ANISA ALIYANI', 'DEDE DIANAH', 'FITRI',
            'HAPIATULIANI', 'INDI NOPIA AYU', 'RESYA NURUL AWALIA', 'REVA YULIA AGUSTI',
            'REVA YULIA AGUSTIN', 'RISNA', 'SELVI PEBRIANTI', 'SINDI AULIA',
            'SITI MUNAWAROH', 'SIVA APRIANI', 'YUYUN YUNITA'
        }

        # 2. Parse Worksheet (Data Akademik, Kehadiran, & Ekskul)
        students_raw = []
        for r in range(8, ws.max_row + 1):
            no_cell = ws.cell(row=r, column=1).value
            nama_cell = ws.cell(row=r, column=2).value
            if not no_cell or not nama_cell:
                continue

            no = int(no_cell)
            nama = str(nama_cell).strip().upper()
            nisn = str(ws.cell(row=r, column=3).value or '').strip()
            nis = str(ws.cell(row=r, column=4).value or '').strip()
            rata_val = ws.cell(row=r, column=21).value
            if rata_val is None:
                continue
            rata = round(float(rata_val), 2)

            sakit = int(ws.cell(row=r, column=22).value or 0)
            izin = int(ws.cell(row=r, column=23).value or 0)
            alpa = int(ws.cell(row=r, column=24).value or 0)
            total_absen = sakit + izin + alpa

            # Hitung persentase kehadiran
            if total_absen == 0:
                persen_hadir = 100.0
            elif total_absen <= 3:
                persen_hadir = 90.0
            elif total_absen <= 15:
                persen_hadir = 80.0
            else:
                persen_hadir = 70.0

            # Hitung keaktifan organisasi dari ekskul (kolom 25-33)
            ekskul_count = sum(1 for c in range(25, 34) if ws.cell(row=r, column=c).value is not None)
            if ekskul_count >= 2:
                keaktifan = 'Aktif'
            elif ekskul_count == 1:
                keaktifan = 'Cukup'
            else:
                keaktifan = 'Tidak'

            # Tentukan jenis kelamin
            jk = 'P' if nama in FEMALE_NAMES else 'L'

            s2_info = sheet2_data.get(nama)
            if not s2_info:
                # fallback partial search
                for k, v in sheet2_data.items():
                    if k in nama or nama in k:
                        s2_info = v
                        break
            if not s2_info:
                s2_info = {
                    'sktm': False,
                    'penghasilan': 2500000,
                    'prestasi': 'Rendah',
                    'jto': 2
                }

            students_raw.append({
                'no': no,
                'nama': nama,
                'nis': nis,
                'nisn': nisn,
                'kelas': 'X',
                'jenis_kelamin': jk,
                'alamat': 'SINDANGRATU',
                'no_hp': '0873673576326',
                'nilai_rata_rata': rata,
                'persentase_hadir': persen_hadir,
                'tingkat_prestasi': s2_info['prestasi'],
                'keaktifan_organisasi': keaktifan,
                'penghasilan_ortu': s2_info['penghasilan'],
                'jumlah_tanggungan': s2_info['jto'],
                'status_rumah': 'Milik Sendiri',
                'memiliki_kip': s2_info['sktm'],
                'jarak_rumah': 'Sedang',
            })

        if not students_raw:
            self.stderr.write(self.style.ERROR("Tidak ada data siswa yang ditemukan di Worksheet."))
            return

        # 3. Hitung Ranking Kelas berdasarkan nilai rata-rata tertinggi
        students_sorted = sorted(students_raw, key=lambda x: x['nilai_rata_rata'], reverse=True)
        for rank, st in enumerate(students_sorted, 1):
            st['ranking_kelas'] = rank

        # 4. Simpan ke Database dalam transaksi atomic
        created_count = 0
        updated_count = 0

        with transaction.atomic():
            for st in sorted(students_raw, key=lambda x: x['no']):
                siswa, s_created = Siswa.objects.update_or_create(
                    nis=st['nis'],
                    defaults={
                        'nama': st['nama'],
                        'kelas': st['kelas'],
                        'jenis_kelamin': st['jenis_kelamin'],
                        'alamat': st['alamat'],
                        'no_hp': st['no_hp'],
                    }
                )

                penilaian, p_created = PenilaianBeasiswa.objects.update_or_create(
                    siswa=siswa,
                    defaults={
                        'nilai_rata_rata': st['nilai_rata_rata'],
                        'ranking_kelas': st['ranking_kelas'],
                        'persentase_hadir': st['persentase_hadir'],
                        'tingkat_prestasi': st['tingkat_prestasi'],
                        'keaktifan_organisasi': st['keaktifan_organisasi'],
                        'penghasilan_ortu': st['penghasilan_ortu'],
                        'jumlah_tanggungan': st['jumlah_tanggungan'],
                        'status_rumah': st['status_rumah'],
                        'memiliki_kip': st['memiliki_kip'],
                        'jarak_rumah': st['jarak_rumah'],
                    }
                )

                if s_created:
                    created_count += 1
                else:
                    updated_count += 1

                self.stdout.write(
                    f"[{'BARU' if s_created else 'UPDATE'}] {st['no']:2}. NIS: {st['nis']} - {st['nama']} | "
                    f"JK: {st['jenis_kelamin']} | Rata: {st['nilai_rata_rata']} (Rank {st['ranking_kelas']}) | "
                    f"KIP: {'Ya' if st['memiliki_kip'] else 'Tidak'} | "
                    f"Gaji: Rp {st['penghasilan_ortu']:,}"
                )

        self.stdout.write(self.style.SUCCESS(
            f"\nSelesai! Berhasil mengimpor {created_count} siswa baru dan memperbarui {updated_count} siswa kelas X."
        ))

        # 5. Opsi Jalankan Seleksi C4.5
        if options['run_seleksi']:
            training_data = DataTraining.objects.all()
            if training_data.count() < 2:
                self.stderr.write(self.style.WARNING("Data training kurang dari 2. Seleksi dilewati."))
                return

            self.stdout.write(self.style.NOTICE("\nMenjalankan perhitungan C4.5 untuk siswa Kelas X..."))
            for st in students_raw:
                siswa = Siswa.objects.get(nis=st['nis'])
                hasil = run_seleksi(siswa.penilaian, training_data)
                HasilSeleksi.objects.update_or_create(
                    penilaian=siswa.penilaian,
                    defaults={
                        'status_seleksi': hasil['status'],
                        'entropy_awal': hasil['entropy_awal'],
                        'root_node': hasil['root_node'],
                        'rules_json': hasil['rules_json'],
                        'gain_info_json': hasil['gain_info_json'],
                        'keterangan': f'Dihitung dengan {training_data.count()} data training.',
                    }
                )
                self.stdout.write(f"  -> {siswa.nama}: {hasil['status']}")

            self.stdout.write(self.style.SUCCESS("Perhitungan C4.5 selesai."))
