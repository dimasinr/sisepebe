import os
import openpyxl
from django.core.management.base import BaseCommand
from django.db import transaction
from beasiswa.models import Siswa, PenilaianBeasiswa, DataTraining, HasilSeleksi
from beasiswa.utils.c45 import run_seleksi


class Command(BaseCommand):
    help = "Import data siswa dan penilaian kelas XI dari file Excel (f_leger_XI A.xlsx / f_leger_XI B.xlsx)"

    def add_arguments(self, parser):
        parser.add_argument(
            '--file',
            type=str,
            default='f_leger_XI B.xlsx',
            help='Path ke file Excel leger nilai kelas XI'
        )
        parser.add_argument(
            '--run-seleksi',
            action='store_true',
            help='Otomatis jalankan klasifikasi C4.5 setelah import'
        )

    def handle(self, *args, **options):
        file_path = options['file']
        if not os.path.exists(file_path):
            self.stderr.write(self.style.ERROR(f"File tidak ditemukan: {file_path}"))
            return

        self.stdout.write(self.style.NOTICE(f"Membaca file: {file_path}..."))

        wb = openpyxl.load_workbook(file_path, data_only=True)
        
        # 1. Tentukan Sheet Nilai (Worksheet) dan Sheet Sosial (Sheet1/Sheet2)
        ws = None
        s_soc = None

        if 'Worksheet' in wb.sheetnames:
            ws = wb['Worksheet']
        else:
            ws = wb.active

        for sname in wb.sheetnames:
            if sname == ws.title:
                continue
            sheet_candidate = wb[sname]
            # Cek apakah sheet ini mengandung data SKTM / Penghasilan
            for r in range(1, min(5, sheet_candidate.max_row + 1)):
                row_vals = [str(sheet_candidate.cell(row=r, column=c).value or '').upper() for c in range(1, sheet_candidate.max_column + 1)]
                if any('SKTM' in v or 'PENGHASILAN' in v for v in row_vals):
                    s_soc = sheet_candidate
                    break
            if s_soc:
                break

        if not s_soc:
            # Fallback ke Sheet2 atau Sheet1
            if 'Sheet2' in wb.sheetnames:
                s_soc = wb['Sheet2']
            elif 'Sheet1' in wb.sheetnames:
                s_soc = wb['Sheet1']
            else:
                self.stderr.write(self.style.ERROR("Sheet kriteria sosial/ekonomi tidak ditemukan."))
                return

        # 2. Parse Sheet Sosial & Ekonomi
        soc_header_row = 1
        for r in range(1, min(5, s_soc.max_row + 1)):
            row_vals = [str(s_soc.cell(row=r, column=c).value or '').upper() for c in range(1, s_soc.max_column + 1)]
            if any('SKTM' in v or 'PENGHASILAN' in v for v in row_vals):
                soc_header_row = r
                break

        sktm_col = 5
        penghasilan_col = 6
        prestasi_col = 7
        jto_col = 8

        for c in range(1, s_soc.max_column + 1):
            val = str(s_soc.cell(row=soc_header_row, column=c).value or '').upper()
            if 'SKTM' in val or 'KIP' in val:
                sktm_col = c
            elif 'PENGHASILAN' in val or 'GAJI' in val:
                penghasilan_col = c
            elif 'PRESTASI' in val:
                prestasi_col = c
            elif 'JTO' in val or 'TANGGUNGAN' in val:
                jto_col = c

        sheet_soc_data = {}
        for r in range(soc_header_row + 1, s_soc.max_row + 1):
            no_cell = s_soc.cell(row=r, column=1).value
            nama_cell = s_soc.cell(row=r, column=2).value
            if not no_cell and not nama_cell:
                continue

            try:
                no = int(no_cell)
            except (TypeError, ValueError):
                no = None

            nama = str(nama_cell or '').strip().upper()
            nisn = str(s_soc.cell(row=r, column=3).value or '').strip()
            nis = str(s_soc.cell(row=r, column=4).value or '').strip()
            sktm_val = str(s_soc.cell(row=r, column=sktm_col).value or '').strip().upper()
            penghasilan_val = s_soc.cell(row=r, column=penghasilan_col).value or 0
            prestasi_val = str(s_soc.cell(row=r, column=prestasi_col).value or 'RENDAH').strip().title()
            jto_val = s_soc.cell(row=r, column=jto_col).value or 1

            entry = {
                'no': no,
                'nama': nama,
                'nis': nis,
                'nisn': nisn,
                'sktm': sktm_val in ['IYA', 'YA', 'TRUE', '1'],
                'penghasilan': int(penghasilan_val),
                'prestasi': prestasi_val if prestasi_val in ['Tinggi', 'Sedang', 'Rendah'] else 'Rendah',
                'jto': int(jto_val),
            }

            if nis:
                sheet_soc_data[nis] = entry
            if no:
                sheet_soc_data[f"no_{no}"] = entry
            if nama:
                sheet_soc_data[f"nama_{nama}"] = entry

        # 3. Cari kolom Rata-rata, Ketidakhadiran, Ekskul di Worksheet
        ws_rata_col = 18
        sakit_col = 19
        izin_col = 20
        alpa_col = 21

        for c in range(1, ws.max_column + 1):
            val4 = str(ws.cell(row=4, column=c).value or '').upper()
            val5 = str(ws.cell(row=5, column=c).value or '').upper()
            if 'RATA' in val4 or 'RATA' in val5:
                ws_rata_col = c
            elif 'SAKIT' in val4 or 'SAKIT' in str(ws.cell(row=7, column=c).value or '').upper():
                sakit_col = c
            elif 'IZIN' in val4 or 'IZIN' in str(ws.cell(row=7, column=c).value or '').upper():
                izin_col = c
            elif 'ALPA' in val4 or 'ALPA' in str(ws.cell(row=7, column=c).value or '').upper():
                alpa_col = c

        ekskul_start_col = max(alpa_col + 1, ws_rata_col + 4)
        ekskul_end_col = ws.max_column

        # Mapping Nama Laki-laki khusus
        MALE_NAMES = {
            'ABIM GIAN CANDRA', 'AGUS RENDI', 'ALPAN APANDI', 'DEDE CAHYANI', 'DEDE HERMAWAN',
            'DEDE NOVA ARYANTO', 'DESTA PEBRIANA', 'GALANG', 'IMAM ARIPIN', 'KENDI',
            'MARLIN', 'MOHAMAD ALI', 'MOHAMAD ARIFIN JULIANSYAH', 'MUHAMAD AKBAR PERDANA',
            'MUHAMAD RIZKI', 'NADIN RAMADANI', 'NURYAMAN', 'PAIZ BAGJAN', 'PAJRIL ABI NOVA',
            'PREDI SUPRIADI', 'PEBRI SUPRIADI', 'ROBIANSYAH', 'SAHRUL GUNAWAN', 'SURYADI',
            'TRI FAHMI FAIRUZ', 'TRIYANA SEPTIAN RAMADHAN', 'WAHID', 'YAMAN SUPRIADI'
        }

        # 4. Parse Data Siswa dari Worksheet
        students_raw = []
        for r in range(8, ws.max_row + 1):
            no_cell = ws.cell(row=r, column=1).value
            nama_cell = ws.cell(row=r, column=2).value
            if not no_cell or not nama_cell:
                continue

            try:
                no = int(no_cell)
            except (TypeError, ValueError):
                continue

            nama = str(nama_cell).strip().upper()
            nisn = str(ws.cell(row=r, column=3).value or '').strip()
            nis = str(ws.cell(row=r, column=4).value or '').strip()
            rata_val = ws.cell(row=r, column=ws_rata_col).value
            if rata_val is None:
                continue
            rata = round(float(rata_val), 2)

            sakit = int(ws.cell(row=r, column=sakit_col).value or 0)
            izin = int(ws.cell(row=r, column=izin_col).value or 0)
            alpa = int(ws.cell(row=r, column=alpa_col).value or 0)
            total_absen = sakit + izin + alpa

            # Hitung persentase kehadiran
            if total_absen == 0:
                persen_hadir = 100.0
            elif total_absen <= 3:
                persen_hadir = 90.0
            elif total_absen <= 15:
                persen_hadir = 80.0
            else:
                persen_hadir = 70.0

            # Hitung keaktifan organisasi dari ekskul
            ekskul_count = sum(1 for c in range(ekskul_start_col, ekskul_end_col + 1) if ws.cell(row=r, column=c).value is not None)
            if ekskul_count >= 2:
                keaktifan = 'Aktif'
            elif ekskul_count == 1:
                keaktifan = 'Cukup'
            else:
                keaktifan = 'Tidak'

            # Tentukan jenis kelamin
            if nama in MALE_NAMES:
                jk = 'L'
            elif len(nis) >= 5 and nis[4] == '1':
                jk = 'L'
            else:
                jk = 'P'

            # Cari data sosial dari sheet kriteria
            s_info = sheet_soc_data.get(nis) or sheet_soc_data.get(f"no_{no}") or sheet_soc_data.get(f"nama_{nama}")
            if not s_info:
                s_info = {
                    'sktm': False,
                    'penghasilan': 2500000,
                    'prestasi': 'Rendah',
                    'jto': 2
                }

            students_raw.append({
                'no': no,
                'nama': nama,
                'nis': nis,
                'nisn': nisn,
                'kelas': 'XI',
                'jenis_kelamin': jk,
                'alamat': 'SINDANGRATU',
                'no_hp': '0873673576326',
                'nilai_rata_rata': rata,
                'persentase_hadir': persen_hadir,
                'tingkat_prestasi': s_info['prestasi'],
                'keaktifan_organisasi': keaktifan,
                'penghasilan_ortu': s_info['penghasilan'],
                'jumlah_tanggungan': s_info['jto'],
                'status_rumah': 'Milik Sendiri',
                'memiliki_kip': s_info['sktm'],
                'jarak_rumah': 'Sedang',
            })

        if not students_raw:
            self.stderr.write(self.style.ERROR("Tidak ada data siswa yang ditemukan di Worksheet."))
            return

        # 5. Hitung Ranking Kelas berdasarkan nilai rata-rata tertinggi
        students_sorted = sorted(students_raw, key=lambda x: x['nilai_rata_rata'], reverse=True)
        for rank, st in enumerate(students_sorted, 1):
            st['ranking_kelas'] = rank

        # 6. Simpan ke Database dalam transaksi atomic
        created_count = 0
        updated_count = 0

        with transaction.atomic():
            for st in sorted(students_raw, key=lambda x: x['no']):
                siswa, s_created = Siswa.objects.update_or_create(
                    nis=st['nis'],
                    defaults={
                        'nama': st['nama'],
                        'kelas': st['kelas'],
                        'jenis_kelamin': st['jenis_kelamin'],
                        'alamat': st['alamat'],
                        'no_hp': st['no_hp'],
                    }
                )

                penilaian, p_created = PenilaianBeasiswa.objects.update_or_create(
                    siswa=siswa,
                    defaults={
                        'nilai_rata_rata': st['nilai_rata_rata'],
                        'ranking_kelas': st['ranking_kelas'],
                        'persentase_hadir': st['persentase_hadir'],
                        'tingkat_prestasi': st['tingkat_prestasi'],
                        'keaktifan_organisasi': st['keaktifan_organisasi'],
                        'penghasilan_ortu': st['penghasilan_ortu'],
                        'jumlah_tanggungan': st['jumlah_tanggungan'],
                        'status_rumah': st['status_rumah'],
                        'memiliki_kip': st['memiliki_kip'],
                        'jarak_rumah': st['jarak_rumah'],
                    }
                )

                if s_created:
                    created_count += 1
                else:
                    updated_count += 1

                self.stdout.write(
                    f"[{'BARU' if s_created else 'UPDATE'}] {st['no']:2}. NIS: {st['nis']} - {st['nama']} | "
                    f"JK: {st['jenis_kelamin']} | Rata: {st['nilai_rata_rata']} (Rank {st['ranking_kelas']}) | "
                    f"KIP: {'Ya' if st['memiliki_kip'] else 'Tidak'} | "
                    f"Gaji: Rp {st['penghasilan_ortu']:,}"
                )

        self.stdout.write(self.style.SUCCESS(
            f"\nSelesai! Berhasil mengimpor {created_count} siswa baru dan memperbarui {updated_count} siswa kelas XI."
        ))

        # 7. Opsi Jalankan Seleksi C4.5
        if options['run_seleksi']:
            training_data = DataTraining.objects.all()
            if training_data.count() < 2:
                self.stderr.write(self.style.WARNING("Data training kurang dari 2. Seleksi dilewati."))
                return

            self.stdout.write(self.style.NOTICE("\nMenjalankan perhitungan C4.5 untuk siswa Kelas XI..."))
            for st in students_raw:
                siswa = Siswa.objects.get(nis=st['nis'])
                hasil = run_seleksi(siswa.penilaian, training_data)
                HasilSeleksi.objects.update_or_create(
                    penilaian=siswa.penilaian,
                    defaults={
                        'status_seleksi': hasil['status'],
                        'entropy_awal': hasil['entropy_awal'],
                        'root_node': hasil['root_node'],
                        'rules_json': hasil['rules_json'],
                        'gain_info_json': hasil['gain_info_json'],
                        'keterangan': f'Dihitung dengan {training_data.count()} data training.',
                    }
                )
                self.stdout.write(f"  -> {siswa.nama}: {hasil['status']}")

            self.stdout.write(self.style.SUCCESS("Perhitungan C4.5 selesai."))
